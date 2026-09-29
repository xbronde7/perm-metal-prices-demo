from __future__ import annotations

import csv
import io
import json
import re
from pathlib import PurePosixPath
from urllib.parse import unquote, urljoin, urlparse
from xml.etree import ElementTree as ET

from bs4 import BeautifulSoup
import openpyxl
import xlrd

from .domain import attributes, identity, number, text, unit

FILE_RE = re.compile(r"\.(csv|xlsx?|pdf|docx|xml)(?:$|[?#])", re.I)
DATE_RE = re.compile(r"\b\d{2}\.\d{2}\.20\d{2}\b")


def decode(blob):
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return blob.decode(enc)
        except UnicodeDecodeError:
            pass
    return blob.decode("utf-8", "replace")


def links(blob, url):
    if blob.lstrip().startswith(b"<?xml"):
        try:
            root=ET.fromstring(blob)
            return [{"url":n.text,"text":n.text,"file":bool(FILE_RE.search(n.text))} for n in root.iter() if n.tag.split("}")[-1]=="loc" and n.text and n.text.startswith(("https://","http://"))][:300]
        except ET.ParseError:
            return []
    soup = BeautifulSoup(blob, "html.parser")
    found = {}
    for a in soup.select("a[href]"):
        href = a.get("href", "")
        if not href or href.startswith(("#", "javascript:", "mailto:", "tel:")):
            continue
        target = urljoin(url, href).split("#")[0]
        if urlparse(target).scheme not in {"http", "https"}:
            continue
        found[target] = {"url":target,"text":text(a.get_text(" "))[:160],"file":bool(FILE_RE.search(unquote(target)))}
    return list(found.values())


def vat_status(context):
    if re.search(r"без\s+НДС", context, re.I):
        return "excluded"
    if re.search(r"(?:с\s+НДС|включая\s+НДС)", context, re.I):
        return "included"
    return "unknown"


def header_map(row, mode):
    out = {}
    for i, raw in enumerate(row):
        value = text(raw).lower()
        if re.search(r"наименован|номенклатур|название|^товар$|^позиция$", value):
            out.setdefault("name", i)
        if re.search(r"размер|характерист|описание|гост|марка\s+стал", value):
            out.setdefault("spec", []).append(i)
        if re.search(r"кол[ -]?во|количест|остат|наличие", value):
            out.setdefault("quantity", i)
        if re.search(r"ед\.?\s*(?:изм|измер)|единица|^ед\.?$", value):
            out.setdefault("unit", i)
        if re.search(r"цена|розничн|стоимост.*(?:единиц|ед\.)", value) and not re.search(r"сумма|общая|стоимость.*всего", value):
            out.setdefault("price", i)
            m = re.search(r"(?:/|за\s+)(тн?|кг|км|м|шт)\b", value)
            if m:
                out["priceUnit"] = unit(m[1])
        if re.search(r"^код$|артикул|sku", value):
            out.setdefault("sku", i)
    necessary = "quantity" if mode == "demand" else "price"
    return out if "name" in out and necessary in out else None


def table_records(rows, url, mode="offer", context=""):
    records, mapping, category = [], None, ""
    context = text(context)
    price_date = DATE_RE.search(context)
    for line_no, row in enumerate(rows):
        values = [text(x) for x in row]
        fresh = header_map(values, mode)
        if fresh:
            mapping = fresh
            context = " ".join(values) + " " + context
            continue
        if not mapping:
            continue
        cell = lambda key: values[mapping[key]] if key in mapping and mapping[key] < len(values) else ""
        name = cell("name")
        if not name or name.lower().startswith(("итого", "всего")):
            continue
        if len([v for v in values if v]) == 1:
            category = name
            continue
        spec = " ".join(values[i] for i in mapping.get("spec", []) if i < len(values) and values[i] != name)
        full_name = text(name + " " + spec)
        if not attributes(full_name, category)["family"]:
            continue
        qty, price = number(cell("quantity")), number(cell("price"))
        if mode == "offer" and (price is None or price <= 0):
            continue
        if mode == "demand" and (qty is None or qty <= 0):
            qty = None
        u = unit(cell("unit")) or mapping.get("priceUnit")
        rec = {"id":identity(url,line_no,full_name),"name":full_name,"category":category,"quantity":qty,"unit":u,"price":price,"sourceUrl":url,"evidence":f"Строка {line_no+1}","vat":vat_status(context),"priceDate":price_date[0] if price_date else None,"sku":cell("sku") or None}
        rec["attributes"] = attributes(full_name,category)
        if mode == "offer":
            rec["stockQuantity"] = qty
            rec["stockUnit"] = u if qty is not None else None
            rec["stockStatus"] = "reported" if qty is not None else "unknown"
        records.append(rec)
    return records


def extract(blob, url, content_type="", mode="offer"):
    ext = PurePosixPath(unquote(urlparse(url).path)).suffix.lower()
    tables, notes, context = [], [], ""
    if blob.startswith(b"PK\x03\x04") and (ext == ".xlsx" or "spreadsheet" in content_type):
        book = openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
        for sheet in book:
            rows = [list(r) for r in sheet.values]
            tables.append((rows," ".join(text(x) for r in rows[:15] for x in r)))
        book.close()
    elif blob.startswith(b"\xd0\xcf\x11\xe0") or ext == ".xls":
        book = xlrd.open_workbook(file_contents=blob)
        for sheet in book.sheets():
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
            tables.append((rows," ".join(text(x) for r in rows[:15] for x in r)))
    elif blob.startswith(b"%PDF"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(blob)) as pdf:
            for page in pdf.pages[:80]:
                page_text = page.extract_text() or ""
                context += page_text[:2000] + "\n"
                for table in page.extract_tables():
                    tables.append((table,page_text[:1800]))
        if not context.strip():
            notes.append("Скан PDF: требуется OCR; значения не извлечены")
        elif not tables:
            notes.append("PDF без распознанной таблицы: требуется настройка формата")
    elif ext == ".docx" or (blob.startswith(b"PK\x03\x04") and "word" in content_type):
        from docx import Document
        doc = Document(io.BytesIO(blob))
        context = " ".join(p.text for p in doc.paragraphs[:40])
        for table in doc.tables:
            tables.append(([[c.text for c in r.cells] for r in table.rows],context))
    elif ext == ".csv" or "text/csv" in content_type:
        raw = decode(blob)
        try:
            dialect = csv.Sniffer().sniff(raw[:5000],delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        rows = list(csv.reader(io.StringIO(raw),dialect))
        tables.append((rows," ".join(text(x) for r in rows[:6] for x in r)))
    elif ext == ".xml" or "xml" in content_type or blob.lstrip().startswith(b"<?xml"):
        root = ET.fromstring(blob)
        records = []
        for o in root.findall(".//offer"):
            name, price = o.findtext("name"), number(o.findtext("price"))
            if not name or not price:
                continue
            u, qty = unit(o.findtext("unit")), number(o.findtext("quantity"))
            records.append({"id":identity(url,o.get("id"),name),"name":text(name),"price":price,"unit":u,"stockQuantity":qty,"stockUnit":u if qty is not None else None,"stockStatus":"reported" if qty is not None else "unknown","vat":"unknown","sourceUrl":url,"attributes":attributes(name),"sku":o.get("id"),"evidence":"XML offer"})
        return records, notes
    else:
        soup = BeautifulSoup(blob,"html.parser")
        title = soup.select_one("h1")
        context = soup.get_text(" ",strip=True)
        for table in soup.select("table"):
            rows = [[c.get_text(" ",strip=True) for c in r.find_all(["td","th"],recursive=False)] for r in table.find_all("tr")]
            tables.append((rows,context[:2000]))
        records = []
        for script in soup.select('script[type="application/ld+json"]'):
            try:
                raw = json.loads(script.string or script.get_text())
            except (ValueError,TypeError):
                continue
            products = raw if isinstance(raw,list) else raw.get("@graph",[raw]) if isinstance(raw,dict) else []
            for obj in products:
                if not isinstance(obj,dict) or obj.get("@type") != "Product":
                    continue
                offers = obj.get("offers") or []
                for offer in offers if isinstance(offers,list) else [offers]:
                    if not isinstance(offer,dict):
                        continue
                    name, price = text(obj.get("name")), number(offer.get("price"))
                    if price and attributes(name)["family"]:
                        m = re.search(r"(?:₽|руб\.?|рублей)\s*(?:/|за)\s*(т|кг|км|м|шт)\b",context,re.I)
                        records.append({"id":identity(url,name,price),"name":name,"price":price,"unit":unit(m[1]) if m else None,"sourceUrl":url,"attributes":attributes(name),"stockQuantity":None,"stockUnit":None,"stockStatus":"listed" if "InStock" in str(offer.get("availability")) else "unknown","vat":"unknown","evidence":"JSON-LD Product"})
        # Many Russian catalogs use schema.org microdata instead of JSON-LD.
        if not records and title:
            p = soup.select_one('[itemprop="price"]')
            val = number(p.get("content") or p.get_text()) if p else None
            name = title.get_text(" ",strip=True)
            if val and attributes(name)["family"]:
                m = re.search(r"(?:₽|руб\.?)\s*(?:/|за)\s*(т|кг|км|м|шт)\b",context,re.I)
                records.append({"id":identity(url,name),"name":name,"price":val,"unit":unit(m[1]) if m else None,"sourceUrl":url,"attributes":attributes(name),"stockQuantity":None,"stockUnit":None,"stockStatus":"listed" if re.search(r"в наличии",context,re.I) else "unknown","vat":"unknown","evidence":"schema.org price"})
        for rows,ctx in tables:
            records.extend(table_records(rows,url,mode,ctx))
        return records, notes
    records = []
    for rows,ctx in tables:
        records.extend(table_records(rows,url,mode,ctx))
    return records, notes
